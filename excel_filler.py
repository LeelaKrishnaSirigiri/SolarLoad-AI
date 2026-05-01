from openpyxl import load_workbook
from datetime import datetime
import os


def clean_number(value):
    if value is None or value == "":
        return ""

    value = str(value)
    value = value.replace("₹", "")
    value = value.replace("Rs.", "")
    value = value.replace("Rs", "")
    value = value.replace(",", "")
    value = value.strip()

    try:
        return float(value)
    except ValueError:
        return value


def fill_excel_template(template_path, bill_data):
    wb = load_workbook(template_path)
    ws = wb.active

    ws["D1"] = bill_data.get("consumer_name", "")
    ws["D2"] = bill_data.get("consumer_number", "")
    ws["D3"] = clean_number(bill_data.get("fixed_charges", ""))
    ws["D4"] = bill_data.get("sanctioned_load_kw", "")
    ws["D5"] = bill_data.get("connection_type", "")
    ws["D7"] = 600

    monthly_units = bill_data.get("monthly_units", {})

    month_row_map = {
        "February 2025": 9,
        "March 2025": 10,
        "April 2025": 11,
        "May 2025": 12,
        "June 2025": 13,
        "July 2025": 14,
        "August 2025": 15,
        "September 2025": 16,
        "October 2025": 17,
        "November 2025": 18,
        "December 2025": 19,
        "January 2026": 20,
    }

    for month, row in month_row_map.items():
        if month in monthly_units and monthly_units[month] != "":
            ws[f"D{row}"] = clean_number(monthly_units[month])

    # Safety: current month units must always go to January row
    ws["D20"] = clean_number(bill_data.get("units_consumed", ""))

    # Bill amount
    ws["E20"] = clean_number(bill_data.get("bill_amount", ""))

    os.makedirs("output", exist_ok=True)

    output_path = f"output/filled_output_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(output_path)

    return output_path
